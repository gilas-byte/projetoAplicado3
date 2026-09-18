#!/usr/bin/env python3
"""Extrai os dicionarios de variaveis do INEP (2019-2022) para JSON.

Atende ao item 5 do Plano de Trabalho ("Contrato de dados": nome INEP,
tipo, dominio/categorias). Os dicionarios vem como XLSX dentro dos zips
originais, num layout em que o nome da variavel aparece so na primeira
linha e as categorias seguem nas linhas abaixo:

    NOME DA VARIAVEL | Descricao   | Categoria | Descricao | Tamanho | Tipo
    TP_FAIXA_ETARIA  | Faixa etaria| 1         | Menor ... | 2       | Numerica
                     |             | 2         | 17 anos   |         |

Alem de converter, o script compara as quatro edicoes e aponta onde o
dominio mudou - que e a lista de trabalho da harmonizacao - e confere a
cobertura contra as colunas realmente presentes nos CSVs.

Saidas (em docs/dicionario/):
  dicionario_<ano>.json - variaveis, tipos e dominios de cada edicao
  README.md             - divergencias entre edicoes e dominios de recorte

Uso:
  ./.venv/bin/python scripts/dicionario.py
"""

from __future__ import annotations

import glob
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))

RAIZ = Path(__file__).resolve().parent.parent
RAW = RAIZ / "data" / "raw"
SAIDA = RAIZ / "docs" / "dicionario"
PROFILING = RAIZ / "docs" / "profiling"

ANOS = [2019, 2020, 2021, 2022]

from colunas import PRESENCA, RECORTE_COM_DOMINIO as RECORTE  # noqa: E402


def achar_xlsx(ano: int) -> Path:
    padrao = str(RAW / f"microdados_enem_{ano}" / "DICION*" / "*.xlsx")
    achados = sorted(glob.glob(padrao))
    if not achados:
        raise FileNotFoundError(
            f"Dicionario de {ano} nao encontrado em {padrao}.\n"
            f"Extraia do zip: unzip -o {RAW / f'microdados_enem_{ano}.zip'} "
            f"'DICION*/*' -d {RAW / f'microdados_enem_{ano}'}/"
        )
    return Path(achados[0])


def texto(valor) -> str:
    """Normaliza celula em string limpa; devolve '' para vazio."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).replace("\n", " ").strip()


def parse_dicionario(ano: int) -> dict:
    """Le a aba de microdados e devolve {variavel: {descricao, tipo, ...}}."""
    caminho = achar_xlsx(ano)
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    aba = next(
        (n for n in wb.sheetnames if n.upper().startswith("MICRODADOS")), wb.sheetnames[0]
    )
    ws = wb[aba]

    variaveis: dict[str, dict] = {}
    atual: str | None = None
    secao = ""

    for linha in ws.iter_rows(values_only=True):
        celulas = [texto(c) for c in linha] + [""] * 6
        nome, descricao, codigo, rotulo, tamanho, tipo = celulas[:6]

        # Cabecalho do arquivo e da tabela.
        if nome.upper().startswith("DICION") or nome.upper().startswith("NOME DA"):
            continue
        # Linha de secao: so a coluna A preenchida (ex.: "DADOS DO PARTICIPANTE").
        if nome and not any([descricao, codigo, rotulo, tamanho, tipo]):
            secao = nome
            continue

        if nome:
            atual = nome
            variaveis[atual] = {
                "variavel": nome,
                "secao": secao,
                "descricao": descricao,
                "tipo": tipo,
                "tamanho": tamanho,
                "categorias": {},
            }

        # Linha de categoria (a primeira vem junto com o nome da variavel,
        # as demais vem com a coluna A vazia).
        if atual and codigo != "":
            variaveis[atual]["categorias"][codigo] = rotulo

    wb.close()
    return {
        "ano": ano,
        "arquivo": str(caminho.relative_to(RAIZ)),
        "aba": aba,
        "variaveis": variaveis,
    }


# ---------------------------------------------------------------- relatorio


def colunas_do_csv(ano: int) -> list[str] | None:
    """Colunas realmente presentes no CSV, segundo o profiling ja gerado."""
    arquivo = PROFILING / f"profiling_{ano}.json"
    if not arquivo.exists():
        return None
    perfil = json.loads(arquivo.read_text(encoding="utf-8"))
    return [c["coluna"] for c in perfil["colunas"]]


def gerar_relatorio(dicts: dict[int, dict]) -> str:
    anos = sorted(dicts)
    L = []
    L.append("# Dicionario de variaveis do ENEM (2019-2022)")
    L.append("")
    L.append(
        "Gerado por `scripts/dicionario.py` a partir dos XLSX oficiais do INEP "
        "contidos nos zips. Dominios completos em `dicionario_<ano>.json`."
    )
    L.append("")

    # --- cobertura -------------------------------------------------------
    L.append("## 1. Cobertura")
    L.append("")
    L.append(
        "Confere o dicionario contra as colunas que o CSV realmente tem "
        "(segundo `docs/profiling/`). Divergencia aqui significa coluna sem "
        "dominio documentado, ou dominio documentado para coluna inexistente."
    )
    L.append("")
    L.append("| Edicao | Vars. no dicionario | Colunas no CSV | So no dicionario | So no CSV |")
    L.append("| --- | --- | --- | --- | --- |")
    for ano in anos:
        vars_dic = set(dicts[ano]["variaveis"])
        cols = colunas_do_csv(ano)
        if cols is None:
            L.append(f"| {ano} | {len(vars_dic)} | (profiling ausente) | - | - |")
            continue
        so_dic = sorted(vars_dic - set(cols))
        so_csv = sorted(set(cols) - vars_dic)
        L.append(
            f"| {ano} | {len(vars_dic)} | {len(cols)} | "
            f"{', '.join(f'`{v}`' for v in so_dic) or '-'} | "
            f"{', '.join(f'`{v}`' for v in so_csv) or '-'} |"
        )
    L.append("")

    todas = sorted({v for ano in anos for v in dicts[ano]["variaveis"]})

    def dominio(var: str, ano: int) -> dict[str, str] | None:
        info = dicts[ano]["variaveis"].get(var)
        return info["categorias"] if info else None

    # --- mesmo codigo, significado diferente -----------------------------
    L.append("## 2. Mesmo codigo, rotulo diferente")
    L.append("")
    L.append(
        "O codigo existe nas quatro edicoes e nenhuma checagem de dominio "
        "acusa problema, mas o rotulo mudou. Nas tres variaveis abaixo o "
        "codigo e relativo a edicao e continua comparavel entre anos; o que "
        "nao se pode e exibir o rotulo de uma edicao como se valesse para "
        "todas. Ver `docs/contrato/` para a regra de cada uma. "
        "`CO_PROVA_*` fica de fora porque o codigo do caderno e reemitido a "
        "cada aplicacao por construcao."
    )
    L.append("")

    remarcados = []
    for var in todas:
        if var.startswith("CO_PROVA_"):
            continue
        por_ano = {a: dominio(var, a) for a in anos if dominio(var, a)}
        if len(por_ano) < 2:
            continue
        comuns = set.intersection(*(set(d) for d in por_ano.values()))
        conflitos = {}
        for codigo in comuns:
            rotulos = {a: por_ano[a][codigo].strip() for a in por_ano}
            if len(set(rotulos.values())) > 1:
                conflitos[codigo] = rotulos
        if conflitos:
            remarcados.append((var, conflitos))

    if not remarcados:
        L.append("Nenhuma variavel remarcou codigos entre as edicoes.")
        L.append("")
    for var, conflitos in remarcados:
        info = next(dicts[a]["variaveis"][var] for a in anos if var in dicts[a]["variaveis"])
        L.append(f"### `{var}` — {info['descricao']}")
        L.append("")
        L.append(f"{len(conflitos)} codigo(s) com rotulo diferente entre edicoes.")
        L.append("")
        L.append("| Codigo | " + " | ".join(str(a) for a in anos) + " |")
        L.append("| --- |" + " --- |" * len(anos))
        for codigo in sorted(conflitos, key=lambda c: (len(c), c)):
            rotulos = conflitos[codigo]
            L.append(
                f"| `{codigo}` | "
                + " | ".join(rotulos.get(a, "-") for a in anos)
                + " |"
            )
        L.append("")

    # --- codigos que entram e saem ---------------------------------------
    L.append("## 3. Codigos que entram ou saem")
    L.append("")
    L.append(
        "O conjunto de codigos nao e o mesmo nas quatro edicoes. Menos "
        "perigoso que o caso anterior, porque um filtro por codigo "
        "inexistente devolve vazio em vez de resposta errada, mas ainda "
        "exige decisao explicita."
    )
    L.append("")

    entram_saem = []
    for var in todas:
        if var.startswith("CO_PROVA_"):
            continue
        por_ano = {a: dominio(var, a) for a in anos if dominio(var, a)}
        if len(por_ano) < 2:
            continue
        conjuntos = [set(d) for d in por_ano.values()]
        if any(c != conjuntos[0] for c in conjuntos):
            entram_saem.append((var, por_ano))

    if not entram_saem:
        L.append("Nenhum.")
        L.append("")
    else:
        L.append("| Variavel | Codigo | Rotulo | Edicoes em que existe |")
        L.append("| --- | --- | --- | --- |")
        for var, por_ano in entram_saem:
            universo = set().union(*(set(d) for d in por_ano.values()))
            comuns = set.intersection(*(set(d) for d in por_ano.values()))
            for codigo in sorted(universo - comuns, key=lambda c: (len(c), c)):
                presentes = [a for a in anos if a in por_ano and codigo in por_ano[a]]
                rotulo = por_ano[presentes[0]][codigo]
                L.append(
                    f"| `{var}` | `{codigo}` | {rotulo} | "
                    f"{', '.join(str(a) for a in presentes)} |"
                )
        L.append("")

    L.append(
        "> `CO_PROVA_CN/CH/LC/MT` foram omitidos das duas secoes acima: o "
        "codigo do caderno de prova e reemitido a cada aplicacao, entao "
        "praticamente nenhum codigo se repete entre edicoes. Os dominios "
        "completos estao nos JSON."
    )
    L.append("")

    # --- documentado x observado -----------------------------------------
    L.append("## 4. Documentado x observado")
    L.append("")
    L.append(
        "Cruza o dominio do dicionario com os valores que aparecem de fato "
        "nos CSVs (`docs/profiling/`). `so no dado` e o caso grave: valor "
        "sem rotulo oficial. `so no dicionario` e categoria prevista que "
        "nunca ocorreu - o filtro existe mas devolveria vazio."
    )
    L.append("")
    L.append("| Variavel | Edicao | So no dado | So no dicionario |")
    L.append("| --- | --- | --- | --- |")
    achou_divergencia = False
    for var in RECORTE:
        for ano in anos:
            perfil_arq = PROFILING / f"profiling_{ano}.json"
            if not perfil_arq.exists():
                continue
            perfil = json.loads(perfil_arq.read_text(encoding="utf-8"))
            if var not in perfil.get("recorte", {}):
                continue
            observados = {
                str(x["categoria"])
                for x in perfil["recorte"][var]
                if x["categoria"] is not None
            }
            documentados = set(dominio(var, ano) or {})
            if not documentados:
                continue
            sem_rotulo = sorted(observados - documentados, key=lambda c: (len(c), c))
            sem_uso = sorted(documentados - observados, key=lambda c: (len(c), c))
            if sem_rotulo or sem_uso:
                achou_divergencia = True
                L.append(
                    f"| `{var}` | {ano} | "
                    + (", ".join(f"`{c}`" for c in sem_rotulo) or "-")
                    + " | "
                    + (", ".join(f"`{c}`" for c in sem_uso) or "-")
                    + " |"
                )
    if not achou_divergencia:
        L.append("| - | - | - | - |")
        L.append("")
        L.append("Nenhuma divergencia: todo valor observado tem rotulo e vice-versa.")
    L.append("")

    # --- dominios de recorte --------------------------------------------
    L.append("## 5. Dominios das variaveis de recorte")
    L.append("")
    L.append(
        "Rotulos da edicao mais recente disponivel. Sao as variaveis que "
        "viram filtro no produto."
    )
    L.append("")
    for var in RECORTE + PRESENCA:
        ano_ref = next((a for a in reversed(anos) if var in dicts[a]["variaveis"]), None)
        if ano_ref is None:
            continue
        info = dicts[ano_ref]["variaveis"][var]
        L.append(f"### `{var}` — {info['descricao']}")
        L.append("")
        L.append(f"Tipo: {info['tipo'] or '-'} · fonte: dicionario {ano_ref}")
        L.append("")
        if not info["categorias"]:
            L.append("Sem dominio categorico no dicionario.")
            L.append("")
            continue
        L.append("| Codigo | Rotulo |")
        L.append("| --- | --- |")
        for codigo, rotulo in info["categorias"].items():
            L.append(f"| `{codigo}` | {rotulo} |")
        L.append("")

    return "\n".join(L)


def main() -> int:
    SAIDA.mkdir(parents=True, exist_ok=True)
    dicts = {}
    for ano in ANOS:
        print(f"[{ano}] lendo dicionario...", flush=True)
        dicts[ano] = parse_dicionario(ano)
        n_vars = len(dicts[ano]["variaveis"])
        n_cats = sum(len(v["categorias"]) for v in dicts[ano]["variaveis"].values())
        print(f"[{ano}] {n_vars} variaveis, {n_cats} categorias", flush=True)
        destino = SAIDA / f"dicionario_{ano}.json"
        destino.write_text(
            json.dumps(dicts[ano], indent=2, ensure_ascii=False), encoding="utf-8"
        )
        print(f"[{ano}] -> {destino.relative_to(RAIZ)}", flush=True)

    relatorio = SAIDA / "README.md"
    relatorio.write_text(gerar_relatorio(dicts), encoding="utf-8")
    print(f"relatorio -> {relatorio.relative_to(RAIZ)}", flush=True)
    return 0


if __name__ == "__main__":
    sys.exit(main())
